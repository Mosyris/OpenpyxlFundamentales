from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Series, Reference


# create table content
wb = Workbook()
ws = wb.active
ws.tittle = "Tree Infos"
treeData = [["Type", "Leaf Color", "Height"], ["Maple", "Red", 549], ["Oak", "Green", 783], ["Pine", "Green", 1204]]
for row in treeData:
    ws.append(row)

# change the header of the table to bold

ft = Font(bold=True)
for row in ws["A1:C1"]:
    for cell in row:
        cell.font = ft

# Build chart skeleton

chart = BarChart()
chart.type = "col"
chart.title = " Tree Height"
chart.y_axis.title = "Height (cm)"
chart.x_axis.title = "Tree Type"
chart.legend = None

data = Reference(ws,
                 min_col=3,
                 min_row=2,
                 max_row=4,
                 max_col=4)

categories = Reference(ws,
                       min_col=1,
                       min_row=2,
                       max_row=4,
                       max_col=1)

chart.add_data(data)
chart.set_categories(categories)

# reference of the chart
ws.add_chart(chart, "E1")

# render
wb.save("treeData.xlsx")





